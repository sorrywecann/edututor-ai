# tutor-service/app/services/rag_utils.py
"""Shared utilities for RAG services."""
import logging
import os
import tempfile
from pathlib import Path
from typing import List

logger = logging.getLogger(__name__)


def split_text(
    text: str,
    chunk_size: int,
    chunk_overlap: int,
    length_fn,
    separators: list,
) -> list:
    """Recursive text splitter shared by all RAG backends."""
    def _split(t: str, seps: List[str]) -> List[str]:
        sep = seps[0] if seps else ""
        parts = t.split(sep) if sep else list(t)
        chunks: List[str] = []
        current: List[str] = []
        current_len = 0
        for part in parts:
            part_with_sep = (part + sep) if sep else part
            part_len = length_fn(part_with_sep)
            if current_len + part_len > chunk_size and current:
                chunks.append(sep.join(current))
                while current and current_len > chunk_overlap:
                    removed = current.pop(0)
                    current_len -= length_fn((removed + sep) if sep else removed)
            current.append(part)
            current_len += part_len
        if current:
            chunks.append(sep.join(current))
        if len(seps) > 1:
            result: List[str] = []
            for chunk in chunks:
                if length_fn(chunk) > chunk_size:
                    result.extend(_split(chunk, seps[1:]))
                else:
                    result.append(chunk)
            return result
        return chunks

    raw = _split(text, separators)
    return [c.strip() for c in raw if c.strip()]


async def extract_text(content: bytes, filename: str, file_type: str) -> str:
    """Extract plain text from a file blob.

    Supports: PDF, DOCX, XLSX, and plain text/UTF-8 fallback.
    """
    suffix = Path(filename).suffix or f".{file_type}"
    with tempfile.NamedTemporaryFile(delete=False, suffix=suffix) as f:
        f.write(content)
        temp_path = f.name
    try:
        if file_type in ("application/pdf", "pdf"):
            # Try extractors in order: pdfplumber (best layout) → pypdf
            # (broader compatibility) → pypdfium2 (handles edge cases like
            # certain locked / oddly-encoded PDFs). Collect each error so
            # the final RuntimeError tells the user *why* every backend
            # failed — surfacing this in the UI is the only way the user
            # learns the PDF is scanned / image-only / password-protected.
            errors: list[str] = []
            try:
                import pdfplumber
                with pdfplumber.open(temp_path) as pdf:
                    text = "\n\n".join(p.extract_text() or "" for p in pdf.pages).strip()
                if text:
                    return text
                errors.append("pdfplumber: 0 chars extracted")
            except Exception as e:
                errors.append(f"pdfplumber: {e}")

            try:
                import pypdf
                reader = pypdf.PdfReader(temp_path)
                if reader.is_encrypted:
                    try:
                        reader.decrypt("")
                    except Exception:
                        pass
                text = "\n\n".join((p.extract_text() or "") for p in reader.pages).strip()
                if text:
                    return text
                errors.append("pypdf: 0 chars extracted")
            except Exception as e:
                errors.append(f"pypdf: {e}")

            try:
                import pypdfium2 as pdfium
                pdf = pdfium.PdfDocument(temp_path)
                parts: list[str] = []
                for page in pdf:
                    tp = page.get_textpage()
                    parts.append(tp.get_text_range() or "")
                    tp.close()
                    page.close()
                pdf.close()
                text = "\n\n".join(parts).strip()
                if text:
                    return text
                errors.append("pypdfium2: 0 chars extracted")
            except Exception as e:
                errors.append(f"pypdfium2: {e}")

            raise RuntimeError(
                f"PDF {filename!r} contains no extractable text — pravdepodobne "
                f"naskenovaný obrázok alebo chránený dokument. "
                f"Skús ho najprv prejsť cez OCR (napr. Adobe Acrobat → "
                f"Rozpoznať text). Detaily: {' · '.join(errors)}"
            )
        if file_type in (
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            "docx",
        ):
            try:
                import docx2txt
                text = docx2txt.process(temp_path).strip()
                if not text:
                    raise ValueError("DOCX extracted no text")
                return text
            except Exception as e:
                raise RuntimeError(f"DOCX extraction failed for {filename!r}: {e}") from e
        if file_type in (
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            "xlsx",
        ):
            try:
                import openpyxl
                wb = openpyxl.load_workbook(temp_path)
                rows = []
                for sheet in wb.worksheets:
                    for row in sheet.iter_rows(values_only=True):
                        row_text = " | ".join(str(c) for c in row if c)
                        if row_text:
                            rows.append(row_text)
                return "\n".join(rows).strip()
            except Exception as e:
                raise RuntimeError(f"XLSX extraction failed for {filename!r}: {e}") from e
        text = content.decode("utf-8", errors="ignore").strip()
        if not text:
            raise ValueError(f"File {filename!r} appears to be empty")
        return text
    finally:
        os.unlink(temp_path)
